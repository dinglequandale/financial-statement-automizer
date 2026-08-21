"""Sheet discovery and layout detection.

Locates the Balance Sheet / Income Statement sheet inside a workbook and works
out, by inspection rather than by hardcoded column letters, where the labels
live, which row carries the period headers, and which columns hold real
year-end values.

The facts in SPEC-PHASE0.md ("Verified file-format facts") describe what we
*expect* to find (label col A, BS values in C/E, IS "YTD <year>" in Z/BA) --
those are used here only as sanity expectations exercised by tests, never as
constants baked into the detection logic itself.
"""

from __future__ import annotations

import re
from dataclasses import dataclass
from datetime import date, datetime

from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.workbook.workbook import Workbook

from fsa.model.schema import StatementType

# Matches the IS's YTD total columns, e.g. "YTD 2021". Deliberately anchored
# (^...$) so it never matches the monthly "January 2020" style headers.
_YTD_RE = re.compile(r"^YTD\s+(\d{4})$", re.IGNORECASE)

# How many rows from the top we're willing to scan while hunting for the
# header row. Real files put it at row 5 (BS) or row 7 (IS); this just needs
# to be generous enough to never miss it.
_HEADER_SCAN_ROWS = 60

# How far below the header row we scan while counting label-column
# candidates. Real sheets are well under this (BS ~90 rows, IS ~290 rows).
_LABEL_SCAN_ROWS = 500

# Safety caps on how many columns we ever iterate, in case a sheet reports an
# inflated max_column due to stray formatting.
_MAX_COLUMNS = 200


class LayoutError(Exception):
    """Raised when analyze_layout cannot confidently identify a sheet's layout.

    Always carries an actionable message (what was scanned, what wasn't
    found) -- callers should never have to guess why detection failed.
    """


@dataclass(frozen=True)
class ValueColumn:
    column: str  # Excel column letter, e.g. "C"
    fiscal_year: int
    period_end: date | None
    header_raw: str | None


@dataclass(frozen=True)
class SheetLayout:
    sheet: str
    statement: StatementType
    label_column: str
    header_row: int
    value_columns: list[ValueColumn]  # ordered left to right
    entity_name: str | None
    period_caption: str | None


def _norm_sheet_name(name: str) -> str:
    return re.sub(r"\s+", " ", name.strip().lower())


def find_statement_sheet(wb: Workbook, statement: StatementType) -> str | None:
    """Locate the sheet holding the given statement, by name.

    Detection order (per SPEC-PHASE0.md):
      BS: exact "Balance Sheet", exact "Balance Sheets", then any sheet whose
          normalized name contains "balance sheet".
      IS: exact "Consolidated IS", then any sheet whose normalized name
          contains "consolidated is" or "income statement".
    """
    names = wb.sheetnames

    if statement is StatementType.BS:
        for exact in ("Balance Sheet", "Balance Sheets"):
            if exact in names:
                return exact
        for n in names:
            if "balance sheet" in _norm_sheet_name(n):
                return n
        return None

    if statement is StatementType.IS:
        if "Consolidated IS" in names:
            return "Consolidated IS"
        for n in names:
            norm = _norm_sheet_name(n)
            if "consolidated is" in norm or "income statement" in norm:
                return n
        return None

    return None


def _clean_str(value: object) -> str | None:
    if value is None:
        return None
    s = str(value).strip()
    return s if s else None


def _parse_date_cell(value: object) -> date | None:
    """Best-effort date extraction from a header cell.

    Real sample files store these as native datetime values (openpyxl gives
    back datetime.datetime for date-formatted cells), so the isinstance path
    is what actually fires. The string fallback exists only for robustness.
    """
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        s = value.strip()
        for fmt in ("%Y-%m-%d", "%m/%d/%Y", "%B %d, %Y", "%b %d, %Y"):
            try:
                return datetime.strptime(s, fmt).date()
            except ValueError:
                continue
    return None


def analyze_layout(wb: Workbook, sheet: str, statement: StatementType) -> SheetLayout:
    """Detect label column, header row and value columns for one sheet.

    Never guesses silently: raises LayoutError with an actionable message if
    it cannot identify a header row, at least one value column, or a label
    column.
    """
    if sheet not in wb.sheetnames:
        raise LayoutError(
            f"Sheet '{sheet}' not found in workbook. Available sheets: "
            f"{', '.join(wb.sheetnames)}"
        )
    ws = wb[sheet]
    max_row = ws.max_row or 0
    max_col = min(ws.max_column or 0, _MAX_COLUMNS)

    if max_row == 0 or max_col == 0:
        raise LayoutError(f"Sheet '{sheet}' appears to be empty; cannot detect layout")

    entity_name = _clean_str(ws.cell(row=1, column=1).value)
    period_caption = _clean_str(ws.cell(row=3, column=1).value)

    header_scan_limit = min(max_row, _HEADER_SCAN_ROWS)

    best_row: int | None = None
    best_cols: list[tuple[int, int, date | None, object]] = []
    # each entry: (col_index, fiscal_year, period_end_or_None, raw_header_value)

    if statement is StatementType.BS:
        for r in range(1, header_scan_limit + 1):
            found: list[tuple[int, int, date | None, object]] = []
            for c in range(1, max_col + 1):
                raw = ws.cell(row=r, column=c).value
                d = _parse_date_cell(raw)
                if d is not None:
                    found.append((c, d.year, d, raw))
            if len(found) > len(best_cols):
                best_row, best_cols = r, found
        if best_row is None or not best_cols:
            raise LayoutError(
                f"Could not detect a header row with date-valued columns in sheet "
                f"'{sheet}' (statement {statement.value}); scanned rows "
                f"1-{header_scan_limit}. Expected a row where value-column headers "
                f"parse as dates."
            )
    elif statement is StatementType.IS:
        for r in range(1, header_scan_limit + 1):
            found = []
            for c in range(1, max_col + 1):
                raw = ws.cell(row=r, column=c).value
                if isinstance(raw, str):
                    m = _YTD_RE.match(raw.strip())
                    if m:
                        found.append((c, int(m.group(1)), None, raw.strip()))
            if len(found) > len(best_cols):
                best_row, best_cols = r, found
        if best_row is None or not best_cols:
            raise LayoutError(
                f"Could not detect a header row matching '^YTD <year>$' in sheet "
                f"'{sheet}' (statement {statement.value}); scanned rows "
                f"1-{header_scan_limit}."
            )
    else:
        raise LayoutError(f"Unsupported statement type: {statement!r}")

    header_row = best_row
    value_columns = [
        ValueColumn(
            column=get_column_letter(c),
            fiscal_year=year,
            period_end=period_end,
            header_raw=(
                raw.isoformat() if hasattr(raw, "isoformat") else _clean_str(raw)
            ),
        )
        for c, year, period_end, raw in best_cols
    ]
    value_columns.sort(key=lambda vc: column_index_from_string(vc.column))

    value_col_indices = {column_index_from_string(vc.column) for vc in value_columns}

    label_scan_limit = min(max_row, header_row + _LABEL_SCAN_ROWS)
    counts: dict[int, int] = {}
    for c in range(1, max_col + 1):
        if c in value_col_indices:
            continue
        n = 0
        for r in range(header_row + 1, label_scan_limit + 1):
            v = ws.cell(row=r, column=c).value
            if isinstance(v, str) and v.strip():
                n += 1
        counts[c] = n

    if not counts or max(counts.values()) == 0:
        raise LayoutError(
            f"Could not detect a label column in sheet '{sheet}' (statement "
            f"{statement.value}): no non-value column has any non-numeric text "
            f"below header row {header_row} (scanned rows "
            f"{header_row + 1}-{label_scan_limit})."
        )
    label_col_idx = max(counts, key=lambda c: counts[c])
    label_column = get_column_letter(label_col_idx)

    return SheetLayout(
        sheet=sheet,
        statement=statement,
        label_column=label_column,
        header_row=header_row,
        value_columns=value_columns,
        entity_name=entity_name,
        period_caption=period_caption,
    )
