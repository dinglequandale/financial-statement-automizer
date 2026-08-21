"""XLSX reader: layout-agnostic sibling to discover.py/extract.py's BS/IS path.

`discover.py` and `extract.py` locate one named sheet, detect its header row
and value columns, and extract exactly that layout. This reader is dumber and
more general on purpose: it captures every visible sheet's rows verbatim, for
sources where we don't yet know (or care) which sheet holds which statement --
that judgment belongs to `interpret.py`, shared with the ODS and PDF readers.
Numeric coercion is delegated to `extract._to_number` so a cell reads the same
number here as it would through the existing path, rather than reimplementing
that parsing.

Label detection mirrors `ods.py`'s dominant-label-column rule rather than
"first non-empty string cell" per row: a sheet can carry incidental text in a
column that isn't where the account labels live (a units note, a stray
comment), and the first-cell rule would misread it as the label. Tallying
which column carries the most text cells across the whole sheet and treating
only that column as the label source is the same fix applied for the same
reason in the ODS reader -- see that module's docstring, trap 3.

Depth recovery is the one place Excel gives us more than ODS does: QuickBooks'
Excel export can carry real indentation via `cell.alignment.indent`, which the
ODS export loses entirely. Where that signal is absent we fall back to leading
whitespace in the label text itself; where neither exists we leave `depth`
`None`, per part, all-or-nothing -- see `_read_sheet`.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from fsa.ingest.extract import _to_number
from fsa.ingest.raw import RawDoc, RawPart, RawRow, UnreadableSource

#: Row cap per sheet -- real sheets top out around 300; this just bounds a
#: pathological file.
_MAX_ROWS = 2000

#: Column cap per sheet, mirroring discover.py's _MAX_COLUMNS safety cap.
_MAX_COLS = 200

_MAX_CAPTIONS = 40


def read(path: Path) -> RawDoc:
    """Read an .xlsx workbook into the universal intermediate.

    One `RawPart` per visible worksheet (hidden/veryHidden sheets, e.g. the
    Google-Sheets-export internal storage tabs seen in sample 2, are skipped).
    Raises `UnreadableSource` only if the workbook cannot be opened at all --
    an empty sheet is not an error.
    """
    try:
        wb = load_workbook(path, data_only=True, read_only=False)
    except Exception as exc:
        raise UnreadableSource(f"cannot open {path.name} as .xlsx: {exc}") from exc

    out = RawDoc(source=path, kind="xlsx")
    for name in wb.sheetnames:
        ws = wb[name]
        if ws.sheet_state != "visible":
            continue
        out.parts.append(_read_sheet(ws))
    return out


def _cell_kind(value: object) -> str | None:
    """Classify one cell's cached value: "text", "numeric", or None (other)."""
    if isinstance(value, str):
        return "text" if value.strip() else None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return "numeric"
    return None


def _dominant_label_column(
    rows_cells: list[list[tuple[int, str, object]]],
) -> int | None:
    """The column with the most non-empty text cells across the sheet.

    Ties resolve to the leftmost column, matching ods.py's rule.
    """
    counts: dict[int, int] = {}
    for row_cells in rows_cells:
        for col, kind, _value in row_cells:
            if kind == "text":
                counts[col] = counts.get(col, 0) + 1
    if not counts:
        return None
    best = max(counts.values())
    return min(col for col, n in counts.items() if n == best)


def _read_sheet(ws) -> RawPart:
    part = RawPart(name=ws.title)
    max_row = min(ws.max_row or 0, _MAX_ROWS)
    max_col = min(ws.max_column or 0, _MAX_COLS)
    if max_row == 0 or max_col == 0:
        return part

    # Pass 1: classify every cell once; tally text cells per column to find
    # the label column (see module docstring).
    rows_cells: list[list[tuple[int, str, object]]] = []
    for r in range(1, max_row + 1):
        row_cells: list[tuple[int, str, object]] = []
        for c in range(1, max_col + 1):
            v = ws.cell(row=r, column=c).value
            kind = _cell_kind(v)
            if kind == "numeric":
                num = _to_number(v)
                if num is not None:
                    row_cells.append((c, "numeric", num))
            elif kind == "text":
                row_cells.append((c, "text", v))
        rows_cells.append(row_cells)

    if not any(rows_cells):
        # openpyxl reports max_row=max_column=1 even for a sheet with no
        # cells set at all, so the max_row/max_col==0 guard above never
        # fires for that case. A sheet with no content anywhere is still
        # an empty sheet, not one blank row.
        return part

    label_col = _dominant_label_column(rows_cells)

    # Pass 2: decide the depth mode for the whole part, from indent on
    # whichever row's label-column cell actually holds text (the amended
    # per-row "label cell").
    label_cells: dict[int, object] = {}
    any_indent = False
    if label_col is not None:
        for r in range(1, max_row + 1):
            cell = ws.cell(row=r, column=label_col)
            v = cell.value
            if isinstance(v, str) and v.strip():
                label_cells[r] = cell
                align = cell.alignment
                if align is not None and align.indent and int(align.indent) > 0:
                    any_indent = True

    use_indent = any_indent
    use_whitespace = False
    if not use_indent and label_cells:
        use_whitespace = any(
            isinstance(cell.value, str) and cell.value.startswith(" ")
            for cell in label_cells.values()
        )

    # Pass 3: build rows.
    for r, row_cells in enumerate(rows_cells, start=1):
        label: str | None = None
        raw_label: str | None = None
        stray: str | None = None
        stray_col: int | None = None
        first_numeric_col: int | None = None
        values: list[float | None] = []

        for c, kind, v in row_cells:
            if kind == "numeric":
                values.append(v)
                if first_numeric_col is None:
                    first_numeric_col = c
            else:  # text
                if c == label_col:
                    raw_label = v
                    label = v.strip()
                elif stray is None:
                    stray, stray_col = v.strip(), c

        if label is not None:
            locator_col = label_col
        elif stray_col is not None:
            locator_col = stray_col
        else:
            locator_col = first_numeric_col
        locator = f"{get_column_letter(locator_col)}{r}" if locator_col else ""

        if use_indent:
            indent = 0
            cell = label_cells.get(r)
            if cell is not None and cell.alignment is not None and cell.alignment.indent:
                indent = int(cell.alignment.indent)
            depth: int | None = indent
        elif use_whitespace:
            if raw_label is not None:
                lead = len(raw_label) - len(raw_label.lstrip(" "))
                depth = lead // 2
            else:
                depth = 0
        else:
            depth = None

        truncated = bool(label) and (label.endswith("...") or label.endswith("…"))

        part.rows.append(
            RawRow(
                index=r,
                label=label,
                values=values,
                depth=depth,
                locator=locator,
                truncated=truncated,
            )
        )

        if label is not None:
            caption = label if not any(x is not None for x in values) else None
        else:
            caption = stray
        if caption and len(part.captions) < _MAX_CAPTIONS:
            part.captions.append(caption)

    return part
