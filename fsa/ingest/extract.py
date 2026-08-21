"""Turn a workbook into the core data model, using discover.py's layouts.

Two entry points:

  extract_workbook(path, statements=(BS, IS))
      One client single-year workbook -> its PRIMARY and COMPARATIVE
      ExtractedColumn objects (one BS pair, one IS pair, typically).

  extract_reference_grid(path, sheet, statement)
      Weaver's hand-built multi-year consolidation sheet -> a ConsolidatedTable,
      for audit mode.

Row classification follows SPEC-PHASE0.md's "Row classification" section,
applied in order:
  1. label empty/whitespace       -> BLANK
  2. worksheet row < header row   -> TITLE
  3. label present, value cell
     empty or non-numeric         -> SECTION_HEADER
  4. is_total_label(label)        -> SUBTOTAL
  5. otherwise                    -> DATA
"""

from __future__ import annotations

import warnings
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from fsa.ingest.discover import SheetLayout, ValueColumn, analyze_layout, find_statement_sheet
from fsa.ingest.normalize import is_total_label, normalize
from fsa.model.schema import (
    AccountRow,
    CellRef,
    ColumnRole,
    ConsolidatedRow,
    ConsolidatedTable,
    ExtractedColumn,
    RowKind,
    StatementType,
)


def _to_number(value: object) -> float | None:
    """Coerce a cell's cached value to a float, or None if it isn't numeric.

    Blank cells (None) become None, never 0.0 -- that distinction is
    load-bearing downstream (e.g. FY2022 "Prepaid Taxes" is genuinely blank).
    """
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        s = value.strip()
        if not s or s in ("-", "--"):
            return None
        if s.startswith("="):
            # Uncached formula string -- shouldn't happen under data_only=True,
            # but guard rather than misreport a formula as a number.
            return None
        neg = s.startswith("(") and s.endswith(")")
        if neg:
            s = s[1:-1]
        s = s.replace("$", "").replace(",", "").strip()
        if s.endswith("%"):
            s = s[:-1]
        try:
            n = float(s)
        except ValueError:
            return None
        return -n if neg else n
    return None


def _load(path: Path):
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        return load_workbook(path, data_only=True)


@dataclass(frozen=True)
class _RowPlan:
    """Classification for one worksheet row, shared across every value column.

    Row *kind* (and the section it inherits) is decided once per row using
    *all* of the sheet's detected value columns jointly -- not just the one
    column currently being extracted. This matters: a row can be genuinely
    blank in one year's column and populated in another (e.g. FY2022's
    "Prepaid Taxes" is blank in the current-year column but has a value in
    the comparative column). Deciding kind from a single column would
    misclassify that row as a SECTION_HEADER for the blank year; deciding it
    from "is every detected value column blank" correctly keeps it DATA, with
    that particular year's value simply None.
    """

    row_index: int
    label: str
    kind: RowKind
    section: str | None


def _plan_rows(ws, layout: SheetLayout) -> list[_RowPlan]:
    label_idx = column_index_from_string(layout.label_column)
    header_row = layout.header_row
    max_row = ws.max_row or 0
    value_indices = [column_index_from_string(vc.column) for vc in layout.value_columns]

    plan: list[_RowPlan] = []
    current_section: str | None = None

    for r in range(1, max_row + 1):
        raw_label_val = ws.cell(row=r, column=label_idx).value
        label = str(raw_label_val).strip() if raw_label_val is not None else ""

        if label == "":
            kind = RowKind.BLANK
            section: str | None = None
        elif r < header_row:
            kind = RowKind.TITLE
            section = None
        else:
            any_numeric = any(
                _to_number(ws.cell(row=r, column=idx).value) is not None
                for idx in value_indices
            )
            if not any_numeric:
                kind = RowKind.SECTION_HEADER
                current_section = label
                section = None
            elif is_total_label(label):
                kind = RowKind.SUBTOTAL
                section = current_section
            else:
                kind = RowKind.DATA
                section = current_section

        plan.append(_RowPlan(row_index=r, label=label, kind=kind, section=section))

    return plan


def _extract_rows(
    ws, layout: SheetLayout, vc: ValueColumn, path: Path, plan: list[_RowPlan]
) -> list[AccountRow]:
    label_idx = column_index_from_string(layout.label_column)
    value_idx = column_index_from_string(vc.column)

    rows: list[AccountRow] = []
    for entry in plan:
        r = entry.row_index
        label_cell = ws.cell(row=r, column=label_idx)
        value_cell = ws.cell(row=r, column=value_idx)

        if entry.kind in (RowKind.BLANK, RowKind.TITLE, RowKind.SECTION_HEADER):
            value: float | None = None
            ref = CellRef(file=path, sheet=layout.sheet, cell=label_cell.coordinate)
        else:  # DATA or SUBTOTAL -- pull this column's own value, blank stays None
            value = _to_number(value_cell.value)
            ref = CellRef(file=path, sheet=layout.sheet, cell=value_cell.coordinate)

        rows.append(
            AccountRow(
                raw_label=entry.label,
                norm_label=normalize(entry.label),
                kind=entry.kind,
                value=value,
                section=entry.section,
                row_index=r,
                ref=ref,
            )
        )

    return rows


def extract_workbook(
    path: Path,
    statements: tuple[StatementType, ...] = (StatementType.BS, StatementType.IS),
) -> list[ExtractedColumn]:
    """One client workbook -> its PRIMARY and COMPARATIVE ExtractedColumns.

    Role assignment: within each statement's detected value columns, the
    highest fiscal year is PRIMARY, all others are COMPARATIVE.
    """
    path = Path(path)
    wb = _load(path)

    columns: list[ExtractedColumn] = []
    for statement in statements:
        sheet = find_statement_sheet(wb, statement)
        if sheet is None:
            # This workbook simply doesn't carry this statement -- not an error,
            # extract_workbook is asked for both by default but callers may pass
            # a subset, or a workbook may only have one of the two.
            continue

        layout = analyze_layout(wb, sheet, statement)
        ws = wb[sheet]
        if not layout.value_columns:
            continue

        plan = _plan_rows(ws, layout)
        primary_year = max(vc.fiscal_year for vc in layout.value_columns)
        for vc in layout.value_columns:
            role = (
                ColumnRole.PRIMARY
                if vc.fiscal_year == primary_year
                else ColumnRole.COMPARATIVE
            )
            rows = _extract_rows(ws, layout, vc, path, plan)
            columns.append(
                ExtractedColumn(
                    statement=statement,
                    fiscal_year=vc.fiscal_year,
                    role=role,
                    source_file=path,
                    source_sheet=layout.sheet,
                    value_column=vc.column,
                    label_column=layout.label_column,
                    period_end=vc.period_end,
                    rows=rows,
                )
            )

    return columns


def extract_reference_grid(
    path: Path, sheet: str, statement: StatementType
) -> ConsolidatedTable:
    """Weaver's multi-year hand-built consolidation -> ConsolidatedTable.

    Unlike a single-year client file, one row here spans several year columns
    at once, so section/subtotal/data classification is based on whether *any*
    year column holds a numeric value for that row (an account can be
    genuinely absent in some years and present in others).
    """
    path = Path(path)
    wb = _load(path)
    layout = analyze_layout(wb, sheet, statement)
    ws = wb[sheet]

    label_idx = column_index_from_string(layout.label_column)
    header_row = layout.header_row
    max_row = ws.max_row or 0

    value_cols = [(vc, column_index_from_string(vc.column)) for vc in layout.value_columns]

    table = ConsolidatedTable(
        statement=statement, years=[vc.fiscal_year for vc, _ in value_cols]
    )

    current_section: str | None = None

    for r in range(1, max_row + 1):
        label_cell = ws.cell(row=r, column=label_idx)
        raw_label_val = label_cell.value
        label = str(raw_label_val).strip() if raw_label_val is not None else ""

        if label == "" or r < header_row:
            continue

        year_values: dict[int, float | None] = {}
        year_refs: dict[int, CellRef] = {}
        any_numeric = False
        for vc, idx in value_cols:
            cell = ws.cell(row=r, column=idx)
            num = _to_number(cell.value)
            year_values[vc.fiscal_year] = num
            year_refs[vc.fiscal_year] = CellRef(
                file=path, sheet=layout.sheet, cell=cell.coordinate
            )
            if num is not None:
                any_numeric = True

        if not any_numeric:
            # Section header row: no year has a value here.
            current_section = label
            continue

        kind = RowKind.SUBTOTAL if is_total_label(label) else RowKind.DATA

        table.rows.append(
            ConsolidatedRow(
                raw_label=label,
                norm_label=normalize(label),
                kind=kind,
                section=current_section,
                values=year_values,
                provenance=year_refs,
            )
        )

    return table
