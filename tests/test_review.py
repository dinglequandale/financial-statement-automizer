"""Round-trip tests for the Excel review surface.

Pure: builds a tiny MappingSet and table in memory, writes a real .xlsx via
openpyxl, edits it the way an analyst would, and reads it back.
"""

from __future__ import annotations

from pathlib import Path

from fsa.ingest.normalize import normalize
from fsa.mapping.template import TemplateBlock, TemplateLine, TemplateSpec
from fsa.model.mapping import Decider, MappingRule, MappingSet, Target, TargetKind
from fsa.model.schema import (
    ConsolidatedRow,
    ConsolidatedTable,
    RowKind,
    StatementType,
)
from fsa.review.sheet import EXCLUDE_TOKEN, export_review, import_review

BS = StatementType.BS


def _spec() -> TemplateSpec:
    lines = [
        TemplateLine(BS, 9, "Cash and Cash Equivalents", False, False, block="Total Current Assets"),
        TemplateLine(BS, 10, "Accounts Receivable", False, False, block="Total Current Assets"),
        TemplateLine(BS, 11, "Other Assets", False, False, block="Total Current Assets"),
    ]
    blocks = [TemplateBlock(BS, "Total Current Assets", 14, 9, 13)]
    return TemplateSpec(source=Path("x.xlsx"), lines=lines, blocks=blocks)


def _row(label, kind=RowKind.DATA, amount=100.0):
    return ConsolidatedRow(
        raw_label=label, norm_label=normalize(label), kind=kind,
        section="Current Assets", values={2025: amount},
    )


def _fixture():
    table = ConsolidatedTable(statement=BS, years=[2025], rows=[
        _row("Operating Account", amount=5_000.0),
        _row("Mystery Account", amount=9_000_000.0),
        _row("Total Current Assets", RowKind.SUBTOTAL, amount=9_005_000.0),
    ])
    ms = MappingSet(statement=BS, rules=[
        MappingRule("Operating Account", normalize("Operating Account"),
                    Target(BS, "Cash and Cash Equivalents"), decided_by=Decider.ALIAS,
                    confidence=0.95),
        MappingRule("Mystery Account", normalize("Mystery Account"),
                    Target(BS, "UNRESOLVED"), decided_by=Decider.UNRESOLVED, confidence=0.0),
        MappingRule("Total Current Assets", normalize("Total Current Assets"),
                    Target(BS, "UNRESOLVED"), decided_by=Decider.UNRESOLVED, confidence=0.0),
    ])
    return table, ms


def _cells(path):
    from openpyxl import load_workbook
    ws = load_workbook(path)["BS"]
    return {str(ws.cell(row=r, column=2).value): r
            for r in range(3, ws.max_row + 1) if ws.cell(row=r, column=2).value}


def test_unresolved_and_material_sorts_above_resolved(tmp_path):
    table, ms = _fixture()
    p = export_review({BS: ms}, {BS: table}, _spec(), tmp_path / "r.xlsx")
    rows = _cells(p)
    assert rows["Mystery Account"] < rows["Operating Account"], "unresolved must sort first"


def test_client_subtotal_sorts_below_detail_despite_being_largest(tmp_path):
    """A subtotal is large by construction; ranking it on amount alone buries
    the real decisions under a screen of `Total ...` rows."""
    table, ms = _fixture()
    p = export_review({BS: ms}, {BS: table}, _spec(), tmp_path / "r.xlsx")
    rows = _cells(p)
    assert rows["Total Current Assets"] > rows["Mystery Account"]


def test_roundtrip_accept_override_and_exclude(tmp_path):
    from openpyxl import load_workbook

    table, ms = _fixture()
    p = export_review({BS: ms}, {BS: table}, _spec(), tmp_path / "r.xlsx")
    wb = load_workbook(p)
    ws = wb["BS"]
    rows = _cells(p)
    ws.cell(row=rows["Mystery Account"], column=4, value="Accounts Receivable")  # fill
    ws.cell(row=rows["Total Current Assets"], column=4, value=EXCLUDE_TOKEN)     # exclude
    wb.save(p)                                                                   # accept the third

    back, stats = import_review(p, {BS: ms}, _spec())
    out = back[BS]
    assert stats.changed == 1 and stats.excluded == 1 and stats.accepted == 1
    assert stats.unresolved_left == 0
    assert all(r.decided_by.is_confirmed for r in out.rules)
    assert out.for_account(normalize("Mystery Account")).target.label == "Accounts Receivable"
    assert {e.norm_account for e in out.exclusions} == {normalize("Total Current Assets")}


def test_roundtrip_never_loses_an_account(tmp_path):
    table, ms = _fixture()
    p = export_review({BS: ms}, {BS: table}, _spec(), tmp_path / "r.xlsx")
    back, _ = import_review(p, {BS: ms}, _spec())
    assert back[BS].covered() == {r.norm_account for r in ms.rules}


def test_blank_target_stays_unresolved_rather_than_being_guessed(tmp_path):
    table, ms = _fixture()
    p = export_review({BS: ms}, {BS: table}, _spec(), tmp_path / "r.xlsx")
    back, stats = import_review(p, {BS: ms}, _spec())
    assert stats.unresolved_left == 2
    unresolved = [r for r in back[BS].rules if r.decided_by is Decider.UNRESOLVED]
    assert {r.client_account for r in unresolved} == {"Mystery Account", "Total Current Assets"}


def test_new_line_name_becomes_an_insert(tmp_path):
    """Typing a name the template lacks means 'make me a new line'."""
    from openpyxl import load_workbook

    table, ms = _fixture()
    ms.rules[1].target = Target(BS, "Other Assets", block="Total Current Assets")
    ms.rules[1].decided_by = Decider.LLM
    ms.rules[1].confidence = 0.6
    p = export_review({BS: ms}, {BS: table}, _spec(), tmp_path / "r.xlsx")
    wb = load_workbook(p)
    ws = wb["BS"]
    ws.cell(row=_cells(p)["Mystery Account"], column=4, value="Crypto Holdings")
    wb.save(p)

    back, _ = import_review(p, {BS: ms}, _spec())
    rule = back[BS].for_account(normalize("Mystery Account"))
    assert rule.target.kind is TargetKind.INSERT
    assert rule.target.label == "Crypto Holdings"
    assert rule.target.block == "Total Current Assets"


def test_an_auto_exclusion_is_visible_and_reversible(tmp_path):
    """`reconcile` excludes every account a mapped rollup covers. A decision the
    analyst cannot see is a decision they cannot reverse, so those rows must
    appear in the sheet -- and overtyping one must map the account after all."""
    from fsa.model.mapping import Exclusion

    table = ConsolidatedTable(
        statement=BS, years=[2025],
        rows=[_row("Petty Cash"), _row("Rent Deposit")],
    )
    ms = MappingSet(statement=BS)
    ms.rules.append(
        MappingRule(
            client_account="Petty Cash", norm_account=normalize("Petty Cash"),
            target=Target(statement=BS, label="Cash and Cash Equivalents"),
            decided_by=Decider.ALIAS, confidence=0.95,
        )
    )
    ms.exclusions.append(
        Exclusion(
            client_account="Rent Deposit", norm_account=normalize("Rent Deposit"),
            reason="already included in 'Total Current Assets', which is mapped",
            decided_by=Decider.STRUCTURE,
        )
    )

    path = export_review({BS: ms}, {BS: table}, _spec(), tmp_path / "r.xlsx")

    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb[BS.value]
    labels = {ws.cell(row=r, column=2).value: r for r in range(3, ws.max_row + 1)}
    assert "Rent Deposit" in labels, "auto-excluded row vanished from the review sheet"
    assert ws.cell(row=labels["Rent Deposit"], column=4).value == EXCLUDE_TOKEN

    # Untouched: the exclusion is accepted and survives the round trip.
    out, stats = import_review(path, {BS: ms}, _spec())
    assert [e.client_account for e in out[BS].exclusions] == ["Rent Deposit"]

    # Overtyped: the analyst reverses it, and it becomes a real mapping.
    ws.cell(row=labels["Rent Deposit"], column=4).value = "Other Assets"
    wb.save(path)
    out2, _ = import_review(path, {BS: ms}, _spec())
    assert out2[BS].exclusions == []
    reinstated = next(r for r in out2[BS].rules if r.client_account == "Rent Deposit")
    assert reinstated.target.label == "Other Assets"
    assert reinstated.decided_by is Decider.HUMAN


def test_two_accounts_sharing_a_display_label_keep_separate_decisions(tmp_path):
    """The client's own statement can print the same label twice.

    Commercial Flooring's balance sheet has two `Total Credit Cards` rows at
    different nesting levels; the consolidator disambiguates them by section,
    but the review sheet shows both under the same visible name. Reading back
    by display label applied one decision to both and lost the other -- silent
    data loss, and exactly the kind that only appears once the pipeline is
    wired end to end.
    """
    a = ConsolidatedRow(
        raw_label="Total Credit Cards", norm_label="total credit cards",
        kind=RowKind.DATA, section="Credit Cards", values={2025: 1_000.0},
    )
    b = ConsolidatedRow(
        raw_label="Total Credit Cards",
        norm_label="current liabilities total credit cards",
        kind=RowKind.DATA, section="Current Liabilities", values={2025: 2_000.0},
    )
    table = ConsolidatedTable(statement=BS, years=[2025], rows=[a, b])

    ms = MappingSet(statement=BS)
    for row, target in ((a, "Cash and Cash Equivalents"), (b, "Accounts Receivable")):
        ms.rules.append(
            MappingRule(
                client_account=row.raw_label,
                norm_account=row.norm_label,
                target=Target(statement=BS, label=target, kind=TargetKind.ASSIGN),
                decided_by=Decider.ALIAS,
                confidence=0.9,
            )
        )

    path = export_review({BS: ms}, {BS: table}, _spec(), tmp_path / "r.xlsx")
    back, stats = import_review(path, {BS: ms}, _spec())

    assert stats.notes == []  # neither row reported as unknown
    by_norm = {r.norm_account: r.target.label for r in back[BS].rules}
    assert by_norm["total credit cards"] == "Cash and Cash Equivalents"
    assert by_norm["current liabilities total credit cards"] == "Accounts Receivable"
