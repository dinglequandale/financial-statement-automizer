"""Unit tests for fsa.ingest.discover, on synthetic in-memory workbooks.

No dependency on the real client files -- everything here is built with
openpyxl.Workbook() directly, matching the observed shapes described in
SPEC-PHASE0.md without hardcoding them into the detector under test.
"""

from __future__ import annotations

from datetime import datetime

import pytest
from openpyxl import Workbook

from fsa.ingest.discover import (
    LayoutError,
    SheetLayout,
    ValueColumn,
    analyze_layout,
    find_statement_sheet,
)
from fsa.model.schema import StatementType


# ---------------------------------------------------------------------------
# find_statement_sheet
# ---------------------------------------------------------------------------

def test_find_bs_exact_name():
    wb = Workbook()
    wb.active.title = "Balance Sheet"
    wb.create_sheet("Other Stuff")
    assert find_statement_sheet(wb, StatementType.BS) == "Balance Sheet"


def test_find_bs_plural_name_weaver_style():
    wb = Workbook()
    wb.active.title = "Balance Sheets"
    assert find_statement_sheet(wb, StatementType.BS) == "Balance Sheets"


def test_find_bs_fuzzy_fallback():
    wb = Workbook()
    wb.active.title = "FY24 - Balance Sheet (draft)"
    assert find_statement_sheet(wb, StatementType.BS) == "FY24 - Balance Sheet (draft)"


def test_find_bs_returns_none_when_absent():
    wb = Workbook()
    wb.active.title = "Trial Balance"
    wb.create_sheet("Worksheet")
    assert find_statement_sheet(wb, StatementType.BS) is None


def test_find_is_exact_name():
    wb = Workbook()
    wb.active.title = "Consolidated IS"
    assert find_statement_sheet(wb, StatementType.IS) == "Consolidated IS"


def test_find_is_prefers_exact_over_other_sheets():
    wb = Workbook()
    wb.active.title = "Houston"
    wb.create_sheet("Consolidated IS")
    wb.create_sheet("Chicago")
    assert find_statement_sheet(wb, StatementType.IS) == "Consolidated IS"


# ---------------------------------------------------------------------------
# analyze_layout -- balance sheet shape
# ---------------------------------------------------------------------------

def _build_bs_sheet(wb, header_dates: dict[str, datetime], label_col="A"):
    """Minimal BS-shaped sheet: entity/title/caption, dated header row 5,
    a few labeled data rows below, with real numeric values under every
    detected header column.
    """
    ws = wb.active
    ws.title = "Balance Sheet"
    ws["A1"] = "TS Distributors, Inc."
    ws["A2"] = "Balance Sheet"
    ws["A3"] = "As of December 31, 2024"

    for col, d in header_dates.items():
        ws[f"{col}5"] = d
    ws["G5"] = "% net Change"  # never a value column: not a date

    rows = [
        ("Assets", None, None),
        ("Current Assets", None, None),
        ("Cash", 1000.0, 900.0),
        ("Accounts Receivable", 2000.0, 1800.0),
        ("Total Current Assets", 3000.0, 2700.0),
    ]
    value_cols = list(header_dates.keys())
    for i, (label, v1, v2) in enumerate(rows):
        r = 6 + i
        ws[f"{label_col}{r}"] = label
        if v1 is not None:
            ws[f"{value_cols[0]}{r}"] = v1
            ws[f"{value_cols[1]}{r}"] = v2
    return ws


def test_analyze_layout_bs_detects_header_row_and_label_column():
    wb = Workbook()
    _build_bs_sheet(wb, {"C": datetime(2024, 12, 31), "E": datetime(2023, 12, 31)})

    layout = analyze_layout(wb, "Balance Sheet", StatementType.BS)

    assert isinstance(layout, SheetLayout)
    assert layout.label_column == "A"
    assert layout.header_row == 5
    assert layout.entity_name == "TS Distributors, Inc."
    assert layout.period_caption == "As of December 31, 2024"


def test_analyze_layout_bs_value_columns_ordered_and_yeared():
    wb = Workbook()
    _build_bs_sheet(wb, {"C": datetime(2024, 12, 31), "E": datetime(2023, 12, 31)})

    layout = analyze_layout(wb, "Balance Sheet", StatementType.BS)

    assert [vc.column for vc in layout.value_columns] == ["C", "E"]
    assert [vc.fiscal_year for vc in layout.value_columns] == [2024, 2023]
    # left-to-right ordering must hold even if we hand dates in reverse order.
    assert layout.value_columns[0].column < layout.value_columns[1].column


def test_analyze_layout_matches_year_only_not_exact_date():
    """FY2024's real header date is 2024-12-30, not -31. Must still be FY2024."""
    wb = Workbook()
    _build_bs_sheet(wb, {"C": datetime(2024, 12, 30), "E": datetime(2023, 12, 31)})

    layout = analyze_layout(wb, "Balance Sheet", StatementType.BS)

    years = [vc.fiscal_year for vc in layout.value_columns]
    assert years == [2024, 2023]


def test_analyze_layout_ignores_non_date_helper_columns():
    """Columns beyond G may hold scratch numeric work; only date-headed columns count."""
    wb = Workbook()
    ws = _build_bs_sheet(wb, {"C": datetime(2024, 12, 31), "E": datetime(2023, 12, 31)})
    # Helper columns I and K: numeric-looking but their row-5 header is not a date.
    ws["I5"] = "scratch"
    ws["K5"] = "scratch2"
    ws["I6"] = 123.45
    ws["K8"] = 67.89

    layout = analyze_layout(wb, "Balance Sheet", StatementType.BS)

    assert [vc.column for vc in layout.value_columns] == ["C", "E"]


def test_analyze_layout_detects_label_column_not_hardcoded():
    """Labels living in column B (not A) must still be found."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Balance Sheet"
    ws["B1"] = "TS Distributors, Inc."
    ws["D5"] = datetime(2024, 12, 31)
    ws["F5"] = datetime(2023, 12, 31)
    ws["B6"] = "Cash"
    ws["D6"] = 1000.0
    ws["F6"] = 900.0
    ws["B7"] = "Total Current Assets"
    ws["D7"] = 1000.0
    ws["F7"] = 900.0

    layout = analyze_layout(wb, "Balance Sheet", StatementType.BS)

    assert layout.label_column == "B"
    assert [vc.column for vc in layout.value_columns] == ["D", "F"]


# ---------------------------------------------------------------------------
# analyze_layout -- income statement shape
# ---------------------------------------------------------------------------

def _build_is_sheet(wb):
    ws = wb.active
    ws.title = "Consolidated IS"
    ws["A1"] = "TS Distributors, Inc."
    ws["A2"] = "Consolidated Income Statement"

    # Monthly columns B..Y (value, %) pairs -- must be ignored.
    months = [
        "January", "February", "March", "April", "May", "June",
        "July", "August", "September", "October", "November", "December",
    ]
    col = 2  # B
    from openpyxl.utils import get_column_letter
    for m in months:
        ws[f"{get_column_letter(col)}7"] = m
        ws[f"{get_column_letter(col + 1)}7"] = "%"
        col += 2

    ws["Z7"] = "YTD 2021"
    ws["AA7"] = "%"
    # Prior-year monthly columns AC..AZ -- must be ignored.
    col = 29  # AC
    for m in months:
        ws[f"{get_column_letter(col)}7"] = f"{m} 2020"
        ws[f"{get_column_letter(col + 1)}7"] = "%"
        col += 2
    ws["BA7"] = "YTD 2020"
    ws["BB7"] = "%"

    ws["A9"] = "REVENUE"
    ws["A10"] = "Sales"
    ws["Z10"] = 68009052.46
    ws["BA10"] = 58799492.87
    ws["A11"] = "Total Revenue"
    ws["Z11"] = 68009052.46
    ws["BA11"] = 58799492.87
    return ws


def test_analyze_layout_is_detects_ytd_columns_via_regex():
    wb = Workbook()
    _build_is_sheet(wb)

    layout = analyze_layout(wb, "Consolidated IS", StatementType.IS)

    assert layout.header_row == 7
    assert layout.label_column == "A"
    assert len(layout.value_columns) == 2
    assert [vc.column for vc in layout.value_columns] == ["Z", "BA"]
    assert [vc.fiscal_year for vc in layout.value_columns] == [2021, 2020]
    assert layout.value_columns[0].header_raw == "YTD 2021"
    assert layout.value_columns[1].header_raw == "YTD 2020"


def test_analyze_layout_is_ignores_monthly_and_percent_columns():
    wb = Workbook()
    _build_is_sheet(wb)

    layout = analyze_layout(wb, "Consolidated IS", StatementType.IS)

    columns = {vc.column for vc in layout.value_columns}
    assert "B" not in columns  # January
    assert "AA" not in columns  # "%"
    assert "AC" not in columns  # January 2020 (monthly prior year)


# ---------------------------------------------------------------------------
# LayoutError
# ---------------------------------------------------------------------------

def test_analyze_layout_raises_when_no_header_row_found():
    wb = Workbook()
    ws = wb.active
    ws.title = "Balance Sheet"
    ws["A1"] = "Nothing useful here"
    ws["A2"] = "No dates, no YTD headers"

    with pytest.raises(LayoutError):
        analyze_layout(wb, "Balance Sheet", StatementType.BS)


def test_analyze_layout_raises_when_no_label_column_found():
    """Header row detected fine, but no column anywhere has label-like text."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Balance Sheet"
    ws["C5"] = datetime(2024, 12, 31)
    ws["E5"] = datetime(2023, 12, 31)
    # Below the header row, every column is purely numeric -- no labels at all.
    for r in range(6, 10):
        ws[f"C{r}"] = 1.0
        ws[f"E{r}"] = 2.0
        ws[f"F{r}"] = 3.0

    with pytest.raises(LayoutError):
        analyze_layout(wb, "Balance Sheet", StatementType.BS)


def test_analyze_layout_raises_actionable_message():
    wb = Workbook()
    ws = wb.active
    ws.title = "Balance Sheet"
    ws["A1"] = "no header here"

    with pytest.raises(LayoutError) as excinfo:
        analyze_layout(wb, "Balance Sheet", StatementType.BS)

    msg = str(excinfo.value)
    assert msg  # non-empty
    assert "Balance Sheet" in msg


def test_analyze_layout_raises_for_missing_sheet():
    wb = Workbook()
    wb.active.title = "Something Else"

    with pytest.raises(LayoutError):
        analyze_layout(wb, "Balance Sheet", StatementType.BS)


# ---------------------------------------------------------------------------
# Weaver-style multi-year reference grid
# ---------------------------------------------------------------------------

def test_analyze_layout_weaver_style_six_year_grid():
    wb = Workbook()
    ws = wb.active
    ws.title = "Balance Sheets"
    ws["A1"] = "TS Distributors, Inc."
    ws["A2"] = "Balance Sheets"

    years = [2020, 2021, 2022, 2023, 2024, 2025]
    from openpyxl.utils import get_column_letter
    for i, y in enumerate(years):
        col = get_column_letter(2 + i)  # B..G
        ws[f"{col}5"] = datetime(y, 12, 1)

    ws["A6"] = "Assets"
    ws["A7"] = "Current Assets"
    ws["A8"] = "Cash"
    for i in range(len(years)):
        col = get_column_letter(2 + i)
        ws[f"{col}8"] = 1000.0 + i

    layout = analyze_layout(wb, "Balance Sheets", StatementType.BS)

    assert layout.header_row == 5
    assert layout.label_column == "A"
    assert [vc.fiscal_year for vc in layout.value_columns] == years
    assert [vc.column for vc in layout.value_columns] == list("BCDEFG")
