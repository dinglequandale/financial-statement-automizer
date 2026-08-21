"""Unit tests for fsa.ingest.readers.ods and fsa.ingest.readers.xlsx.

Entirely synthetic: .xlsx fixtures are built with openpyxl directly; .ods
fixtures are built by hand-assembling a minimal content.xml (only the parts
the reader actually looks at -- table/row/cell -- inside a real namespaced
office:document-content envelope) and zipping it up. Neither reader looks at
anything else in the archive, so a real .ods's mimetype/META-INF/styles.xml
scaffolding is unnecessary here.
"""

from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.styles import Alignment

from fsa.ingest.raw import UnreadableSource
from fsa.ingest.readers import ods, xlsx

# ---------------------------------------------------------------------------
# .ods fixture helpers
# ---------------------------------------------------------------------------

_CONTENT_TEMPLATE = """<?xml version="1.0" encoding="UTF-8"?>
<office:document-content
xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:1.0"
xmlns:table="urn:oasis:names:tc:opendocument:xmlns:table:1.0"
xmlns:text="urn:oasis:names:tc:opendocument:xmlns:text:1.0"
office:version="1.1">
<office:body>
<office:spreadsheet>
{tables}
</office:spreadsheet>
</office:body>
</office:document-content>
"""


def _text_cell(s: str, repeat: int | None = None) -> str:
    rep = f' table:number-columns-repeated="{repeat}"' if repeat else ""
    return f'<table:table-cell office:value-type="string"{rep}><text:p>{s}</text:p></table:table-cell>'


def _empty_cell(repeat: int | None = None) -> str:
    rep = f' table:number-columns-repeated="{repeat}"' if repeat else ""
    return f'<table:table-cell office:value-type="string"{rep}><text:p></text:p></table:table-cell>'


def _num_cell(value: float, display: str | None = None, repeat: int | None = None) -> str:
    disp = display if display is not None else value
    rep = f' table:number-columns-repeated="{repeat}"' if repeat else ""
    return (
        f'<table:table-cell office:value-type="float" office:value="{value}"{rep}>'
        f"<text:p>{disp}</text:p></table:table-cell>"
    )


def _row(cells: list[str], repeat: int | None = None) -> str:
    rep = f' table:number-rows-repeated="{repeat}"' if repeat else ""
    return f"<table:table-row{rep}>{''.join(cells)}</table:table-row>"


def _table(name: str, rows: list[str]) -> str:
    return f'<table:table table:name="{name}">{"".join(rows)}</table:table>'


def _make_ods(tmp_path: Path, *tables: str, filename: str = "test.ods") -> Path:
    content = _CONTENT_TEMPLATE.format(tables="".join(tables))
    path = tmp_path / filename
    with zipfile.ZipFile(path, "w") as zf:
        zf.writestr("content.xml", content)
    return path


# ---------------------------------------------------------------------------
# ods.py
# ---------------------------------------------------------------------------


def test_ods_label_value_split(tmp_path):
    table = _table(
        "Sheet1",
        [
            _row([_empty_cell(), _text_cell("Total Current Assets"), _num_cell(1990148.32)]),
        ],
    )
    doc = ods.read(_make_ods(tmp_path, table))
    assert doc.kind == "ods"
    assert len(doc.parts) == 1
    row = doc.parts[0].rows[0]
    assert row.label == "Total Current Assets"
    assert row.values == [1990148.32]


def test_ods_numeric_value_attr_preferred_over_display_text(tmp_path):
    # office:value carries the exact number; the paragraph text is a rounded
    # display string that must never be parsed.
    table = _table(
        "Sheet1",
        [_row([_text_cell("Total"), _num_cell(1990148.32, display="1,990,148")])],
    )
    doc = ods.read(_make_ods(tmp_path, table))
    assert doc.parts[0].rows[0].values == [1990148.32]


def test_ods_row_repeat_expansion(tmp_path):
    table = _table(
        "Sheet1",
        [
            _row([_text_cell("Foo"), _num_cell(10)]),
            _row([_empty_cell(), _empty_cell()], repeat=3),
            _row([_text_cell("Bar"), _num_cell(20)]),
        ],
    )
    doc = ods.read(_make_ods(tmp_path, table))
    rows = doc.parts[0].rows
    # Foo, 3 expanded blanks, Bar == 5 rows; row numbers advance through the
    # repeat block so Bar lands on row 5, not row 3.
    assert len(rows) == 5
    assert rows[0].label == "Foo"
    assert rows[0].locator == "A1"
    assert rows[-1].label == "Bar"
    assert rows[-1].locator == "A5"
    assert rows[-1].index == 5


def test_ods_row_repeat_above_cap_is_clamped(tmp_path):
    # A repeat count this large is ODS's nominal 1,048,576-row padding, not
    # real content; expanding it literally would allocate gigabytes.
    table = _table(
        "Sheet1",
        [
            _row([_text_cell("Foo"), _num_cell(10)]),
            _row([_empty_cell(), _empty_cell()], repeat=1048576),
            _row([_text_cell("Bar"), _num_cell(20)]),
        ],
    )
    doc = ods.read(_make_ods(tmp_path, table))
    rows = doc.parts[0].rows
    # The huge-repeat row is clamped to a single row, not expanded.
    assert len(rows) == 3
    assert rows[-1].label == "Bar"
    assert rows[-1].locator == "A3"


def test_ods_column_repeat_above_cap_is_clamped(tmp_path):
    table = _table(
        "Sheet1",
        [_row([_text_cell("Foo"), _num_cell(10, repeat=1024)])],
    )
    doc = ods.read(_make_ods(tmp_path, table))
    row = doc.parts[0].rows[0]
    # Clamped to a single numeric cell, not 1024 copies of the value.
    assert row.values == [10.0]


def test_ods_dominant_label_column_with_furniture_to_the_left(tmp_path):
    # Mirrors the real sample: column A carries sparse print furniture (a
    # timestamp, a date), column B carries the actual body labels. The first
    # non-empty-cell rule would misread "09/23/25" as a label. Row 4 mirrors
    # the real file's "Aug 31, 25" column-header row, where B is blank and
    # the only text sits in column C -- the case where stray text becomes a
    # caption because there is no label-column text to prefer over it.
    table = _table(
        "Sheet1",
        [
            _row([_text_cell("11:12 AM"), _text_cell("Commercial Flooring, Inc.")]),
            _row([_text_cell("09/23/25"), _text_cell("Balance Sheet")]),
            _row([_text_cell("Accrual Basis"), _text_cell("As of August 31, 2025")]),
            _row([_empty_cell(), _empty_cell(), _text_cell("Aug 31, 25")]),
            _row([_empty_cell(), _text_cell("ASSETS")]),
            _row([_empty_cell(), _text_cell("Current Assets")]),
        ],
    )
    doc = ods.read(_make_ods(tmp_path, table))
    rows = doc.parts[0].rows
    labels = [r.label for r in rows]
    assert labels == [
        "Commercial Flooring, Inc.",
        "Balance Sheet",
        "As of August 31, 2025",
        None,
        "ASSETS",
        "Current Assets",
    ]
    # The furniture text in column A never becomes a label...
    assert "09/23/25" not in labels
    assert "11:12 AM" not in labels
    # A row whose label column is blank surfaces its stray text as a caption
    # instead of being silently dropped.
    assert "Aug 31, 25" in doc.parts[0].captions
    # But furniture in a row that DOES have a real label is simply discarded
    # -- only the empty-label-column case promotes stray text to a caption.
    assert "09/23/25" not in doc.parts[0].captions
    assert "11:12 AM" not in doc.parts[0].captions


def test_ods_dominant_label_column_tie_breaks_leftmost(tmp_path):
    table = _table(
        "Sheet1",
        [
            _row([_text_cell("A"), _empty_cell()]),
            _row([_empty_cell(), _text_cell("B")]),
        ],
    )
    doc = ods.read(_make_ods(tmp_path, table))
    rows = doc.parts[0].rows
    # Column 1 and column 2 each have exactly one text cell (a tie) --
    # resolves to the leftmost column, so "B" (in column 2) is furniture.
    assert rows[0].label == "A"
    assert rows[1].label is None
    assert "B" in doc.parts[0].captions


def test_ods_truncated_flag(tmp_path):
    table = _table(
        "Sheet1",
        [
            _row([_text_cell("Accounts Receivable..."), _num_cell(1)]),
            _row([_text_cell("Accounts Payable…"), _num_cell(2)]),
            _row([_text_cell("Cash"), _num_cell(3)]),
        ],
    )
    doc = ods.read(_make_ods(tmp_path, table))
    rows = doc.parts[0].rows
    assert rows[0].truncated is True
    assert rows[1].truncated is True
    assert rows[2].truncated is False


def test_ods_empty_sheet(tmp_path):
    table = _table("Empty", [])
    doc = ods.read(_make_ods(tmp_path, table))
    assert len(doc.parts) == 1
    assert doc.parts[0].name == "Empty"
    assert doc.parts[0].rows == []
    assert doc.parts[0].captions == []


def test_ods_depth_is_always_none(tmp_path):
    table = _table("Sheet1", [_row([_text_cell("Cash"), _num_cell(1)])])
    doc = ods.read(_make_ods(tmp_path, table))
    assert doc.parts[0].rows[0].depth is None
    assert doc.has_depth is False


def test_ods_unreadable_not_a_zip(tmp_path):
    path = tmp_path / "bad.ods"
    path.write_bytes(b"not a zip file at all")
    with pytest.raises(UnreadableSource):
        ods.read(path)


# ---------------------------------------------------------------------------
# xlsx.py
# ---------------------------------------------------------------------------


def _save(wb: Workbook, tmp_path: Path, name: str = "test.xlsx") -> Path:
    path = tmp_path / name
    wb.save(path)
    return path


def test_xlsx_label_value_split(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Balance Sheet"
    ws["A1"] = "Total Current Assets"
    ws["C1"] = 2603316.91
    doc = xlsx.read(_save(wb, tmp_path))
    assert doc.kind == "xlsx"
    part = doc.parts[0]
    assert part.name == "Balance Sheet"
    row = part.rows[0]
    assert row.label == "Total Current Assets"
    assert row.values == [2603316.91]
    assert row.locator == "A1"


def test_xlsx_dominant_label_column_with_furniture_to_the_left(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    # Column A: sparse furniture. Column B: the real labels (dominant).
    ws["A1"], ws["B1"] = "printed 9/23/25", "Balance Sheet"
    ws["B2"] = "ASSETS"
    ws["C3"] = "Aug 31, 25"  # label column B is blank in this row
    ws["B4"] = "Current Assets"
    doc = xlsx.read(_save(wb, tmp_path))
    rows = doc.parts[0].rows
    assert rows[0].label == "Balance Sheet"
    assert rows[1].label == "ASSETS"
    assert rows[2].label is None
    assert rows[3].label == "Current Assets"
    assert "printed 9/23/25" not in [r.label for r in rows]
    # Furniture in a row that also has a real label is discarded, not
    # captioned; a row whose label column is blank promotes its stray text.
    assert "printed 9/23/25" not in doc.parts[0].captions
    assert "Aug 31, 25" in doc.parts[0].captions


def test_xlsx_indent_sets_depth_for_whole_part(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "Current Assets"
    ws["A2"] = "Cash"
    ws["A2"].alignment = Alignment(indent=1)
    ws["B2"] = 100
    ws["A3"] = "Total Current Assets"
    ws["B3"] = 100
    doc = xlsx.read(_save(wb, tmp_path))
    rows = doc.parts[0].rows
    # One row carries a nonzero indent -> every row in the part gets an int
    # depth, unset ones default to 0 (never a mix of int and None).
    assert [r.depth for r in rows] == [0, 1, 0]
    assert all(isinstance(r.depth, int) for r in rows)


def test_xlsx_no_indent_falls_back_to_leading_whitespace(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "Current Assets"
    ws["A2"] = "  Cash"
    ws["B2"] = 100
    ws["A3"] = "    Petty Cash"
    ws["B3"] = 5
    doc = xlsx.read(_save(wb, tmp_path))
    rows = doc.parts[0].rows
    assert rows[0].depth == 0
    assert rows[1].depth == 1  # 2 leading spaces // 2
    assert rows[2].depth == 2  # 4 leading spaces // 2
    # Labels are stripped even though indentation was measured before that.
    assert rows[1].label == "Cash"
    assert rows[2].label == "Petty Cash"


def test_xlsx_no_signal_leaves_depth_none(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "Current Assets"
    ws["A2"] = "Cash"
    ws["B2"] = 100
    doc = xlsx.read(_save(wb, tmp_path))
    rows = doc.parts[0].rows
    assert all(r.depth is None for r in rows)
    assert doc.has_depth is False


def test_xlsx_truncated_flag(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "Accounts Receivable..."
    ws["B1"] = 1
    ws["A2"] = "Accounts Payable…"
    ws["B2"] = 2
    ws["A3"] = "Cash"
    ws["B3"] = 3
    doc = xlsx.read(_save(wb, tmp_path))
    rows = doc.parts[0].rows
    assert rows[0].truncated is True
    assert rows[1].truncated is True
    assert rows[2].truncated is False


def test_xlsx_empty_sheet(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Empty"
    doc = xlsx.read(_save(wb, tmp_path))
    assert len(doc.parts) == 1
    assert doc.parts[0].rows == []
    assert doc.parts[0].captions == []


def test_xlsx_skips_non_visible_sheets(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Visible"
    ws["A1"] = "Cash"
    ws["B1"] = 1
    hidden = wb.create_sheet("Hidden")
    hidden.sheet_state = "hidden"
    hidden["A1"] = "should not appear"
    very_hidden = wb.create_sheet("VeryHidden")
    very_hidden.sheet_state = "veryHidden"
    doc = xlsx.read(_save(wb, tmp_path))
    assert [p.name for p in doc.parts] == ["Visible"]


def test_xlsx_unreadable_not_a_workbook(tmp_path):
    path = tmp_path / "bad.xlsx"
    path.write_bytes(b"not a real xlsx file")
    with pytest.raises(UnreadableSource):
        xlsx.read(path)
